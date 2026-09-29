import io
import json
from typing import Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_json(data: Dict[str, Any]) -> bytes:
    """Return JSON formatted bytes of the analysis result."""
    return json.dumps(data, indent=2, ensure_ascii=False).encode("utf-8")


def export_excel(data: Dict[str, Any]) -> bytes:
    """
    Generate an openpyxl multi-sheet workbook with sheets:
    - Metadata
    - Opinion
    - Exceptions
    - Subservice Orgs
    - CUECs & Mapping
    All with bold header row and auto-sized column widths.
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    header_font = Font(bold=True)

    def style_and_autosize(ws):
        # Make row 1 bold
        for cell in ws[1]:
            cell.font = header_font

        # Auto-size columns based on maximum cell length
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    lines = val_str.split("\n")
                    line_len = max(len(l) for l in lines)
                    max_len = max(max_len, line_len)
                else:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 70)

    # 1. Metadata Sheet
    ws_meta = wb.create_sheet(title="Metadata")
    ws_meta.append(["Field", "Value"])
    meta = data.get("metadata", {})
    ws_meta.append(["Service Organization", meta.get("org", "")])
    ws_meta.append(["System in Scope", meta.get("system", "")])
    ws_meta.append(["Report Type", meta.get("report_type", "")])
    ws_meta.append(["Period Start", meta.get("period_start", "")])
    ws_meta.append(["Period End", meta.get("period_end", "")])
    ws_meta.append(["Independent Auditor", meta.get("auditor", "")])
    ws_meta.append(["Trust Services Criteria", ", ".join(meta.get("criteria", []))])
    style_and_autosize(ws_meta)

    # 2. Opinion Sheet
    ws_op = wb.create_sheet(title="Opinion")
    ws_op.append(["Opinion Type", "Qualifications", "Emphasis of Matter"])
    op = data.get("opinion", {})
    quals = "\n".join(op.get("qualifications", [])) or "None"
    emph = "\n".join(op.get("emphasis_of_matter", [])) or "None"
    ws_op.append([op.get("type", "Unqualified"), quals, emph])
    style_and_autosize(ws_op)

    # 3. Exceptions Sheet
    ws_exc = wb.create_sheet(title="Exceptions")
    ws_exc.append(["Control ID", "Criteria", "Category", "Description", "What Failed (Details)", "Management Response"])
    exceptions = data.get("exceptions", [])
    if exceptions:
        for exc in exceptions:
            ws_exc.append([
                exc.get("control_id", ""),
                exc.get("criteria", ""),
                exc.get("category", ""),
                exc.get("description", ""),
                exc.get("details", exc.get("result", "")),
                exc.get("management_response", "Not provided")
            ])
    else:
        ws_exc.append(["None", "", "", "No control exceptions found.", "", ""])
    style_and_autosize(ws_exc)

    # 4. Subservice Orgs Sheet
    ws_sub = wb.create_sheet(title="Subservice Orgs")
    ws_sub.append(["Subservice Organization", "Reporting Method", "Services Provided", "CSOCs", "Risk Flag"])
    subs = data.get("subservice_orgs", [])
    if subs:
        for s in subs:
            csocs_str = "\n".join(s.get("csocs", []))
            ws_sub.append([
                s.get("name", ""),
                s.get("method", "Carve-Out"),
                s.get("services", ""),
                csocs_str,
                s.get("risk_flag", "HIGH")
            ])
    else:
        ws_sub.append(["None", "", "No subservice organizations disclosed.", "", ""])
    style_and_autosize(ws_sub)

    # 5. CUECs & Mapping Sheet
    ws_cuec = wb.create_sheet(title="CUECs & Mapping")
    ws_cuec.append(["CUEC ID", "Category", "CUEC Text", "Mapped Internal Control", "Similarity Score", "Status"])
    cuecs = data.get("cuecs", [])
    if cuecs:
        for c in cuecs:
            ws_cuec.append([
                c.get("id", ""),
                c.get("category", ""),
                c.get("text", ""),
                c.get("mapped_control", "No matching control"),
                c.get("score", 0.0),
                c.get("status", "Gap")
            ])
    else:
        ws_cuec.append(["None", "", "No CUECs detected.", "", 0.0, ""])
    style_and_autosize(ws_cuec)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
